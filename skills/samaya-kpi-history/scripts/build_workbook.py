"""Build the samaya-kpi-history workbook from a JSON spec.

Usage: python build_workbook.py spec.json out.xlsx   (requires openpyxl)

Spec:
{
  "title": "XPO North American LTL — YoY % per day",
  "metrics": [{"name": "Tonnage/day YoY %", "format": "pct"}, ...],   # format: pct | number
  "rows": [
    {"company": "XPO", "period": "Jan 2023", "status": "Final", "note": "",
     "values": {"Tonnage/day YoY %": {"value": 2.8, "title": "...", "date": "2023-05-04",
                                       "url": "https://app.samaya.ai/u/...", "excerpt": "..."}}}
  ],
  "notes": ["Disclosure map and breaks, one line each"],
  "scorecard": {"Periods found": "44/48", ...}
}
"""

import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

GAP_STATUSES = {"Not disclosed", "Not retrieved", "Not yet reported"}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
GAP_FILL = PatternFill("solid", fgColor="FCE4D6")
LINK_FONT = Font(color="0563C1", underline="single")
FORMATS = {"pct": '0.0"%";(0.0"%")', "number": "#,##0.00"}


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL


def note_text(src: dict) -> str:
    parts = [f"{src['title']} ({src['date']}), via Samaya."]
    if src.get("excerpt"):
        parts.append(f"\"{src['excerpt']}\"")
    parts.append(src["url"])
    return "\n".join(parts)


def build(spec: dict, out: str) -> None:
    metrics = [m["name"] for m in spec["metrics"]]
    fmt = {m["name"]: FORMATS.get(m.get("format", "number")) for m in spec["metrics"]}
    multi = len({r.get("company", "") for r in spec["rows"]}) > 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Series"
    lead = ["Company", "Period"] if multi else ["Period"]
    ws.append(lead + metrics + ["Status"] + [f"Source: {m}" for m in metrics] + ["Note"])
    style_header(ws)
    first_value_col = len(lead) + 1

    for r in spec["rows"]:
        vals = r.get("values", {})
        lead_vals = [r.get("company", ""), r["period"]] if multi else [r["period"]]
        ws.append(
            lead_vals
            + [vals.get(m, {}).get("value") for m in metrics]
            + [r["status"]]
            + [vals.get(m, {}).get("title", "") for m in metrics]
            + [r.get("note", "")]
        )
        row = ws.max_row
        for i, m in enumerate(metrics):
            src = vals.get(m)
            cell = ws.cell(row=row, column=first_value_col + i)
            if fmt[m]:
                cell.number_format = fmt[m]
            if src and src.get("value") is not None and src.get("url"):
                cell.comment = Comment(note_text(src), "samaya-kpi-history", width=340, height=110)
                link = ws.cell(row=row, column=first_value_col + len(metrics) + 1 + i)
                link.hyperlink = src["url"]
                link.font = LINK_FONT
        if r["status"] in GAP_STATUSES:
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = GAP_FILL

    for col in ws.columns:
        letter = col[0].column_letter
        header = str(col[0].value)
        ws.column_dimensions[letter].width = 44 if header.startswith("Source") else 60 if header == "Note" else 16
    ws.freeze_panes = ws.cell(row=2, column=first_value_col)

    if not multi and len(spec["rows"]) > 2:
        chart = LineChart()
        chart.title = spec.get("title", "")
        chart.height, chart.width = 8, 22
        chart.add_data(
            Reference(ws, min_col=first_value_col, max_col=first_value_col + len(metrics) - 1, min_row=1, max_row=ws.max_row),
            titles_from_data=True,
        )
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        wb.create_sheet("Chart").add_chart(chart, "A1")

    if multi:
        pv = wb.create_sheet("Pivot")
        companies = sorted({r["company"] for r in spec["rows"]})
        periods = list(dict.fromkeys(r["period"] for r in spec["rows"]))
        cols = [(c, m) for c in companies for m in metrics]
        pv.append(["Period"] + [f"{c} — {m}" for c, m in cols])
        style_header(pv)
        index = {(r["company"], r["period"]): r.get("values", {}) for r in spec["rows"]}
        for p in periods:
            pv.append([p] + [index.get((c, p), {}).get(m, {}).get("value") for c, m in cols])

    ns = wb.create_sheet("Notes & breaks")
    for line in spec.get("notes", []):
        ns.append([line])
    ns.column_dimensions["A"].width = 130
    for cell in ns["A"]:
        cell.alignment = Alignment(wrap_text=True)

    sc = wb.create_sheet("Scorecard")
    for k, v in spec.get("scorecard", {}).items():
        sc.append([k, v])
    sc.column_dimensions["A"].width = 26
    sc.column_dimensions["B"].width = 90
    for cell in sc["A"]:
        cell.font = Font(bold=True)

    wb.save(out)


if __name__ == "__main__":
    build(json.load(open(sys.argv[1])), sys.argv[2])
    print("saved", sys.argv[2])
