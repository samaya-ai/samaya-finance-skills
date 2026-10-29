#!/usr/bin/env python3
"""
Append citation notes to the speaker notes of flagged slides in the ORIGINAL PowerPoint file.

Usage:
    python annotate_pptx_notes.py <original.pptx> <data.json> <output.pptx>

PPTX has no native comment API reachable from python-pptx/pptxgenjs/Office.js in this
environment (confirmed while building samaya-industry-primer), so this writes citations into each
flagged slide's speaker notes instead -- non-destructive to visible slide content. Multiple
items on the same slide are appended as a block, not overwritten one after another.

Only items with location.type == "pptx" are processed. Only run this after the user has
explicitly opted in to inline annotation.
"""
import json
import sys
from collections import defaultdict
from pptx import Presentation

HEADER = "\n--- Fact-Check Notes ---"


def short_note(item):
    """Note text trimmed for a margin pane: no trailing period, lower-cased opener."""
    note = (item.get("note") or "").strip().rstrip(".")
    if not note:
        return ""
    return note[0].lower() + note[1:]


def note_line(item):
    """One short line per item: flag, description, values only where they differ, source."""
    flag = item.get("flag", "")
    desc = item.get("description", "")
    original = item.get("original_value", "")
    sourced = item.get("samaya_value", "")
    note = short_note(item)

    if flag == "Unable to verify":
        return f"{flag} — {desc}: {note or 'no Samaya route for this claim'}"

    if flag == "Match":
        line = f"{flag} — {desc}: {sourced or original}"
    else:
        line = f"{flag} — {desc}: doc {original} | now {sourced}"
    if note:
        line += f"; {note}"

    doc = item.get("source_doc") or "Not disclosed"
    date = item.get("as_of_date") or ""
    url = item.get("source_url") or "no link returned"
    return f"{line}. {doc}{f' ({date})' if date else ''} {url}"


def annotate(src_path, data, out_path):
    prs = Presentation(src_path)
    by_slide = defaultdict(list)
    skipped = []

    for item in data.get("items", []):
        loc = item.get("location") or {}
        if loc.get("type") != "pptx":
            continue
        slide_num = loc.get("slide")
        if not slide_num or slide_num < 1 or slide_num > len(prs.slides):
            skipped.append(item.get("id"))
            continue
        by_slide[slide_num].append(item)

    applied = 0
    for slide_num, items in by_slide.items():
        slide = prs.slides[slide_num - 1]
        notes_slide = slide.notes_slide  # creates one if it doesn't exist
        text_frame = notes_slide.notes_text_frame
        existing = text_frame.text or ""
        block = HEADER + "\n" + "\n".join(note_line(i) for i in items)
        text_frame.text = existing + block
        applied += len(items)

    prs.save(out_path)
    print(f"Wrote {out_path}: {applied} note(s) applied across {len(by_slide)} slide(s).")
    if skipped:
        print(f"Skipped (invalid/missing slide number): {skipped}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[2]) as f:
        data = json.load(f)
    annotate(sys.argv[1], data, sys.argv[3])
