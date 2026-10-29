#!/usr/bin/env python3
"""
Build the samaya-fact-checker audit log workbook from a JSON payload.

Usage:
    python build_audit_log.py <data.json> <audit_log.xlsx>

Input JSON shape (see references/audit-schema.md for full field notes):

{
  "source_file": "Q2_Investor_Deck.pptx",
  "items": [
    {
      "id": "1", "description": "...", "original_value": "...",
      "location": {"type": "pptx", "slide": 4, "detail": "..."},
      "samaya_value": "...", "source_doc": "...", "source_url": "...",
      "as_of_date": "...", "flag": "Match", "note": "..."
    }
  ]
}

This workbook is the audit log itself (the deliverable), not the annotated original file —
see annotate_docx / annotate_xlsx / annotate_pdf / annotate_pptx_notes for writing citations
back into the source document.
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")
BOLD_FONT = Font(name="Arial", bold=True)

FLAG_FILL = {
    "Match": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
    "Stale": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
    "Mismatch": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "Unable to verify": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}

VALID_FLAGS = set(FLAG_FILL.keys())

COLUMNS = [
    ("ID", 6),
    ("Item Description", 30),
    ("Original Stated Value", 22),
    ("Location", 24),
    ("Samaya-Sourced Value", 22),
    ("Source Document", 26),
    ("Source URL", 40),
    ("As-Of Date", 12),
    ("Flag", 16),
    ("Note", 36),
]


def _location_text(loc):
    if not loc:
        return "Not specified"
    t = loc.get("type")
    if t == "pptx":
        base = f"Slide {loc.get('slide', '?')}"
    elif t == "xlsx":
        base = f"{loc.get('sheet', '?')}!{loc.get('cell', '?')}"
    elif t == "pdf":
        base = f"Page {loc.get('page', '?')}"
    elif t == "docx":
        base = loc.get("detail", "Word document")
    else:
        base = t or "Unknown"
    detail = loc.get("detail")
    if detail and t != "docx":
        return f"{base} — {detail}"
    return base


def build(data, out_path):
    items = data.get("items", [])
    source_file = data.get("source_file", "Not specified")

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Log"

    ws["A1"] = "Fact-Check Audit Log"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = f"Source file: {source_file}"
    ws["A2"].font = BODY_FONT
    ws["A3"] = f"Items checked: {len(items)}"
    ws["A3"].font = BODY_FONT

    header_row = 5
    for col_idx, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row = header_row + 1
    flag_counts = {f: 0 for f in VALID_FLAGS}
    for item in items:
        flag = item.get("flag")
        if flag not in VALID_FLAGS:
            raise ValueError(
                f"Item {item.get('id')} has invalid flag {flag!r}; "
                f"must be one of {sorted(VALID_FLAGS)}"
            )
        flag_counts[flag] += 1

        values = [
            item.get("id", ""),
            item.get("description", ""),
            item.get("original_value", ""),
            _location_text(item.get("location")),
            item.get("samaya_value") or "—",
            item.get("source_doc") or "—",
            item.get("source_url") or "—",
            item.get("as_of_date") or "—",
            flag,
            item.get("note", ""),
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 9:  # Flag column
                cell.fill = FLAG_FILL[flag]
                cell.font = BOLD_FONT
        row += 1

    ws.freeze_panes = f"A{header_row + 1}"

    # Summary sheet
    summary = wb.create_sheet("Summary")
    summary["A1"] = "Fact-Check Summary"
    summary["A1"].font = Font(name="Arial", bold=True, size=14)
    summary["A2"] = f"Source file: {source_file}"
    summary["A3"] = f"Total items checked: {len(items)}"
    summary_row = 5
    summary.cell(row=summary_row, column=1, value="Flag").font = HEADER_FONT
    summary.cell(row=summary_row, column=1).fill = HEADER_FILL
    summary.cell(row=summary_row, column=2, value="Count").font = HEADER_FONT
    summary.cell(row=summary_row, column=2).fill = HEADER_FILL
    for i, flag in enumerate(["Match", "Stale", "Mismatch", "Unable to verify"], start=1):
        r = summary_row + i
        summary.cell(row=r, column=1, value=flag).font = BODY_FONT
        summary.cell(row=r, column=1).fill = FLAG_FILL[flag]
        summary.cell(row=r, column=2, value=flag_counts[flag]).font = BODY_FONT
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 10

    wb.save(out_path)
    print(f"Wrote {out_path}: {len(items)} items "
          f"({', '.join(f'{k}={v}' for k, v in flag_counts.items())})")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1]) as f:
        data = json.load(f)
    build(data, sys.argv[2])
